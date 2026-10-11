import asyncio
import hashlib
import hmac
import html
import math
import os
import re
import shutil
import sys
import time
import zipfile
from datetime import datetime, date, time as dtime
from io import BytesIO
from pathlib import Path
from tempfile import SpooledTemporaryFile

import numpy as np
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    import psutil
except Exception:
    psutil = None

try:
    import resource  # Linux/Unix (Streamlit Cloud OK)
except Exception:
    resource = None


# =====================================================================
# CẤU HÌNH CHUNG
# =====================================================================
st.set_page_config(
    page_title="Công cụ chấm công",
    page_icon="🗂️",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# Cache Chromium của pyppeteer vào HOME để ổn định trên Streamlit Cloud
os.environ.setdefault("PYPPETEER_HOME", str(Path.home() / ".cache" / "pyppeteer"))

DEFAULT_WATERMARK = "Liên hệ Nguyễn Huệ Hr ( 0356 227 868 ) - timvieclam.9phut.com"
DEFAULT_THRESHOLD = 8.0

DEFAULT_RENAME_RULES = """Lương giờ 100% = Lương giờ ca ngày
Lương giờ 130% = Lương giờ ca đêm
TC 150% = T/ca ngày
TC 200% = T/ca Đêm
TC ngày CN 200% = T/ca giờ ngày CN
TC đêm CN 270% = T/Ca giờ đêm CN
"""

MODE_1 = "📂 Tách file chấm công"
MODE_2 = "🖼️ Xuất bảng công ra ảnh"
MODE_HINT = {
    MODE_1: "Từ file chấm công gốc, tạo 1 file Excel mới trong đó mỗi nhân viên có 1 sheet riêng.",
    MODE_2: "Từ file Excel tổng hợp, xuất mỗi sheet thành 1 ảnh PNG và nén tất cả thành 1 file ZIP.",
}


# =====================================================================
# GIAO DIỆN (CSS + THÀNH PHẦN DÙNG CHUNG)
# =====================================================================
APP_CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&display=swap');

:root{
  --ink:#14232B; --muted:#5B6B73; --line:#DDE4E8; --bg:#F4F6F7; --surface:#FFFFFF;
  --brand:#0F766E; --brand-dark:#0B5A54; --brand-soft:#E3F3F1;
  --warn:#92400E; --warn-soft:#FEF3C7;
}

.stApp, .stApp p, .stApp label, .stApp li, .stApp h1, .stApp h2, .stApp h3, .stApp h4,
.stApp button, .stApp input, .stApp textarea, .stApp .stMarkdown{
  font-family:'Be Vietnam Pro', system-ui, -apple-system, 'Segoe UI', Roboto, Arial, sans-serif;
}
.stApp{ background:var(--bg); color:var(--ink); }
.block-container{ max-width:1080px; padding-top:1.6rem; padding-bottom:4rem; }

/* Ẩn phần thừa của Streamlit */
#MainMenu, footer{ visibility:hidden; }
header[data-testid="stHeader"]{ background:transparent; }
section[data-testid="stSidebar"], [data-testid="collapsedControl"]{ display:none; }

/* Tiêu đề */
.hero{ margin:0 0 1.1rem 0; }
.hero-title{ font-size:2rem; font-weight:800; letter-spacing:-.02em; line-height:1.15; color:var(--ink); }
.hero-sub{ margin-top:.4rem; color:var(--muted); font-size:1rem; max-width:62ch; line-height:1.55; }

/* Nút chuyển chức năng (radio dạng tab) */
div[data-testid="stRadio"] div[role="radiogroup"]{
  display:inline-flex; flex-wrap:wrap; gap:.25rem; background:#E5EBEE;
  padding:.3rem; border-radius:14px;
}
div[data-testid="stRadio"] div[role="radiogroup"] > label{
  margin:0; padding:.55rem 1.15rem; border-radius:11px; cursor:pointer; transition:background .15s;
}
div[data-testid="stRadio"] div[role="radiogroup"] > label > div:first-child{ display:none; }
div[data-testid="stRadio"] div[role="radiogroup"] > label p{ font-weight:600; font-size:.97rem; color:var(--muted); }
div[data-testid="stRadio"] div[role="radiogroup"] > label:has(input:checked){
  background:var(--surface); box-shadow:0 1px 3px rgba(20,35,43,.14);
}
div[data-testid="stRadio"] div[role="radiogroup"] > label:has(input:checked) p{ color:var(--ink); }
.mode-hint{ margin:.6rem 0 0 .2rem; color:var(--muted); font-size:.95rem; }

/* Các bước */
.step{ display:flex; gap:.8rem; align-items:flex-start; margin:2rem 0 .7rem 0; }
.step-n{
  flex:none; width:1.9rem; height:1.9rem; border-radius:50%; background:var(--ink); color:#fff;
  font-weight:700; font-size:.92rem; display:flex; align-items:center; justify-content:center; margin-top:.1rem;
}
.step-t{ font-weight:700; font-size:1.2rem; line-height:1.3; color:var(--ink); }
.step-d{ color:var(--muted); font-size:.93rem; margin-top:.1rem; }

/* Ô thống kê */
.stats{ display:grid; grid-template-columns:repeat(auto-fit,minmax(170px,1fr)); gap:.75rem; margin:.4rem 0 1rem 0; }
.stat{ background:var(--surface); border:1px solid var(--line); border-radius:14px; padding:.9rem 1.1rem; }
.stat-v{ font-size:1.7rem; font-weight:800; letter-spacing:-.02em; color:var(--ink); line-height:1.1; }
.stat-l{ color:var(--muted); font-size:.88rem; margin-top:.25rem; }
.stat.good{ background:var(--brand-soft); border-color:#BFE3DE; }
.stat.good .stat-v{ color:var(--brand-dark); }
.stat.warn{ background:var(--warn-soft); border-color:#F6DE8D; }
.stat.warn .stat-v{ color:var(--warn); }

/* Khung tải file */
section[data-testid="stFileUploaderDropzone"]{
  border:2px dashed #8FC5BF; background:#FFFFFF; border-radius:16px; padding:1.4rem 1.2rem;
}
section[data-testid="stFileUploaderDropzone"]:hover{ border-color:var(--brand); background:var(--brand-soft); }
[data-testid="stFileUploaderDropzoneInstructions"] span{ display:none; }
[data-testid="stFileUploaderDropzoneInstructions"]::before{
  content:"Kéo thả file Excel (.xlsx) vào đây"; font-weight:600; font-size:1rem; color:var(--ink);
}
section[data-testid="stFileUploaderDropzone"] button{ font-size:0; border-radius:10px; }
section[data-testid="stFileUploaderDropzone"] button::after{ content:"Chọn file"; font-size:.92rem; font-weight:600; }

/* Nút bấm */
.stButton > button, .stDownloadButton > button{
  border-radius:11px; font-weight:600; padding:.55rem 1.2rem; border:1px solid var(--line);
}
button[data-testid="baseButton-primary"]{
  background:var(--brand); border-color:var(--brand); color:#fff; padding:.7rem 1.5rem; font-size:1rem;
}
button[data-testid="baseButton-primary"]:hover{ background:var(--brand-dark); border-color:var(--brand-dark); color:#fff; }

/* Khung gập, thông báo */
div[data-testid="stExpander"]{ background:var(--surface); border:1px solid var(--line); border-radius:14px; }
div[data-testid="stExpander"] summary{ font-weight:600; }
div[data-testid="stAlert"]{ border-radius:12px; }
div[data-testid="stVerticalBlockBorderWrapper"]{ border-radius:14px; }

.footer-note{ margin-top:3rem; color:#8A989F; font-size:.85rem; text-align:center; }

@media (max-width:640px){
  .hero-title{ font-size:1.6rem; }
  .block-container{ padding-left:1rem; padding-right:1rem; }
}
</style>
"""


def inject_css():
    st.markdown(APP_CSS, unsafe_allow_html=True)


def hero():
    st.markdown(
        '<div class="hero"><div class="hero-title">Công cụ chấm công</div>'
        '<div class="hero-sub">Tách bảng chấm công theo từng nhân viên, sau đó xuất mỗi bảng thành ảnh để gửi đi.</div></div>',
        unsafe_allow_html=True,
    )


def step_header(n: int, title: str, desc: str = ""):
    d = f'<div class="step-d">{html.escape(desc)}</div>' if desc else ""
    st.markdown(
        f'<div class="step"><div class="step-n">{n}</div>'
        f'<div><div class="step-t">{html.escape(title)}</div>{d}</div></div>',
        unsafe_allow_html=True,
    )


def stat_cards(items):
    """items: list of (nhãn, giá trị, tone) — tone: '', 'good', 'warn'"""
    cells = "".join(
        f'<div class="stat {tone}"><div class="stat-v">{html.escape(str(val))}</div>'
        f'<div class="stat-l">{html.escape(label)}</div></div>'
        for label, val, tone in items
    )
    st.markdown(f'<div class="stats">{cells}</div>', unsafe_allow_html=True)


# =====================================================================
# CHỨC NĂNG 1: TÁCH FILE CHẤM CÔNG
# =====================================================================
def safe_excel_value(val):
    if pd.isna(val) or val is None:
        return ""
    if isinstance(val, (float, np.floating)):
        return float(round(val, 2))
    if isinstance(val, (np.integer, int)):
        return int(val)
    return val


def to_hhmm(val):
    if pd.isna(val) or val is None or str(val).strip() == "":
        return ""
    if isinstance(val, (float, np.floating)) and 0 <= val < 1:
        total_minutes = int(round(val * 24 * 60))
        h, m = divmod(total_minutes, 60)
        return f"{int(h):02}:{int(m):02}"
    if hasattr(val, "hour") and hasattr(val, "minute"):
        return f"{val.hour:02}:{val.minute:02}"
    val_str = str(val).strip()
    if ":" in val_str:
        parts = val_str.split(":")
        if len(parts) >= 2:
            return f"{int(parts[0]):02}:{int(parts[1]):02}"
    return val_str


def get_header_row_height(header, width=8):
    lines = []
    for cell in header:
        text = str(cell.value) if cell.value else ""
        line_len = max(1, int((width - 1) * 1.5))
        lines.append(math.ceil(len(text) / line_len))
    return max(24, max(lines) * 15)


def weekday_vn(date_val):
    try:
        d = pd.to_datetime(date_val, dayfirst=True)
        weekday_map = {0: "Thứ 2", 1: "Thứ 3", 2: "Thứ 4", 3: "Thứ 5",
                       4: "Thứ 6", 5: "Thứ 7", 6: "Chủ nhật"}
        return weekday_map[d.weekday()]
    except Exception:
        return ""


@st.cache_data(show_spinner=False, max_entries=3)
def load_attendance(file_bytes: bytes):
    """Đọc các sheet 'CT...' và chuẩn hóa. Lỗi dữ liệu -> ValueError (thông báo tiếng Việt)."""
    xls = pd.ExcelFile(BytesIO(file_bytes))
    ct_sheets = [s for s in xls.sheet_names if s.upper().startswith("CT")]
    if not ct_sheets:
        raise ValueError("Không tìm thấy sheet nào có tên bắt đầu bằng “CT” trong file Excel.")

    df = pd.concat(
        [pd.read_excel(xls, sheet_name=s, header=5) for s in ct_sheets],
        ignore_index=True,
    )

    # Thêm cột Thứ (trước cột Ngày)
    if "Ngày" in df.columns and "Thứ" not in df.columns:
        df.insert(df.columns.get_loc("Ngày"), "Thứ", df["Ngày"].apply(weekday_vn))

    # Chuẩn hóa cột giờ về hh:mm
    for col in [c for c in df.columns if any(k in str(c) for k in ["Vào", "Ra"])]:
        df[col] = df[col].apply(to_hhmm)

    # Làm tròn các cột số
    for col in df.select_dtypes(include=[np.number]).columns:
        df[col] = df[col].apply(lambda x: round(x, 2) if pd.notna(x) else x)

    vao = next((c for c in df.columns if "Vào lần 1" in str(c)), None)
    ra = next((c for c in df.columns if "Ra lần 2" in str(c)), None)
    if vao is None or ra is None:
        raise ValueError("Không tìm thấy cột “Vào lần 1” hoặc “Ra lần 2”. Kiểm tra lại dòng tiêu đề (dòng 6) của sheet CT.")
    if "Mã NV" not in df.columns or "Họ tên" not in df.columns:
        raise ValueError("Thiếu cột “Mã NV” hoặc “Họ tên”. Kiểm tra lại dòng tiêu đề (dòng 6) của sheet CT.")

    return df, ct_sheets, df.columns.get_loc(vao), df.columns.get_loc(ra)


def build_employee_workbook(df, idx_vao_lan_1, idx_ra_lan_2, on_progress=None):
    """Trả về (bytes, số nhân viên, số sheet có ô trống) hoặc None nếu không có dữ liệu."""
    groupby_obj = list(df.groupby(["Mã NV", "Họ tên"]))
    total_nv = max(len(groupby_obj), 1)
    count_nv = 0
    count_tab_vangnhat = 0

    pale_yellow_fill = PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid")
    total_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
    header_fill = PatternFill(start_color="FF8C1A", end_color="FF8C1A", fill_type="solid")

    wb_new = openpyxl.Workbook()
    wb_new.remove(wb_new.active)

    for i, ((ma_nv, ho_ten), group) in enumerate(groupby_obj, start=1):
        if on_progress:
            on_progress(i / total_nv, f"Đang xử lý {i}/{total_nv}: {ma_nv} - {ho_ten}")

        region = group.iloc[:, idx_vao_lan_1:]
        arr = pd.Series(region.values.ravel()).dropna().astype(str).str.strip()
        arr = arr[~arr.isin(["", "nan", "NaT", "None"])]
        if arr.empty:
            continue

        count_nv += 1
        sheet_name = f"{ma_nv}_{ho_ten}".replace(" ", "_").replace("/", "_")[:31]
        ws_nv = wb_new.create_sheet(title=sheet_name)
        ws_nv.append([str(c) for c in group.columns])
        for row in group.itertuples(index=False):
            ws_nv.append([safe_excel_value(cell) for cell in row])

        # Dòng tổng
        total_row = []
        for col in group.columns:
            if group[col].dtype in [np.float64, np.int64, float, int]:
                tong = group[col].sum()
                total_row.append(round(tong, 2) if group[col].notna().sum() > 0 else "")
            else:
                total_row.append("TỔNG" if col == group.columns[0] else "")
        ws_nv.append(total_row)
        for cell in ws_nv[ws_nv.max_row]:
            cell.fill = total_fill
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

        # Header
        header_row = ws_nv[1]
        for cell in header_row:
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

        ws_nv.freeze_panes = "A2"

        for ci, column_cells in enumerate(ws_nv.columns):
            letter = column_cells[0].column_letter
            if ci >= idx_vao_lan_1:
                ws_nv.column_dimensions[letter].width = 8
            else:
                length = max(len(str(c.value) if c.value else "") for c in column_cells)
                ws_nv.column_dimensions[letter].width = min(length + 2, 35)
        ws_nv.row_dimensions[1].height = get_header_row_height(header_row, width=8)

        # Tô vàng nhạt ô trống trong vùng Vào lần 1 -> Ra lần 2
        co_boi = False
        for idx, row in enumerate(ws_nv.iter_rows(min_row=2, max_row=ws_nv.max_row - 1)):
            row_data = group.iloc[idx]
            region_row = row_data.iloc[idx_vao_lan_1:idx_ra_lan_2 + 1]
            values_in_region = [
                v for v in region_row
                if not (pd.isna(v) or str(v).strip() in ["", "nan", "NaT", "None"])
            ]
            if len(values_in_region) > 1:
                for offset, value in enumerate(region_row):
                    cell = ws_nv.cell(row=idx + 2, column=idx_vao_lan_1 + 1 + offset)
                    if pd.isna(value) or str(value).strip() in ["", "nan", "NaT", "None"]:
                        cell.fill = pale_yellow_fill
                        co_boi = True
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="center")

        if co_boi:
            ws_nv.sheet_properties.tabColor = "FFFACD"
            count_tab_vangnhat += 1

    if count_nv == 0:
        return None

    out = BytesIO()
    wb_new.save(out)
    return out.getvalue(), count_nv, count_tab_vangnhat


def goto_mode_2():
    st.session_state["mode"] = MODE_2


def page_tach_file():
    step_header(1, "Tải file chấm công gốc", "File Excel (.xlsx) có các sheet tên bắt đầu bằng “CT”.")
    up = st.file_uploader(
        "File chấm công gốc", type=["xlsx"], key="tach_uploader", label_visibility="collapsed"
    )
    if up is None:
        st.info("Chưa có file. Kéo thả file vào khung phía trên để bắt đầu.")
        return

    file_bytes = up.getvalue()
    md5 = hashlib.md5(file_bytes).hexdigest()

    try:
        with st.spinner("Đang đọc file..."):
            df, ct_sheets, idx_vao, idx_ra = load_attendance(file_bytes)
    except ValueError as e:
        st.error(str(e))
        return
    except Exception as e:
        st.error(f"Không đọc được file. Hãy kiểm tra đây có phải file .xlsx hợp lệ không.\n\nChi tiết: {e}")
        return

    n_emp = df[["Mã NV", "Họ tên"]].dropna().drop_duplicates().shape[0]
    stat_cards([
        ("sheet CT được đọc", len(ct_sheets), ""),
        ("dòng chấm công", f"{len(df):,}", ""),
        ("nhân viên", n_emp, ""),
    ])
    with st.expander("Xem trước dữ liệu (30 dòng đầu)"):
        st.dataframe(df.head(30).astype(str), use_container_width=True, hide_index=True)

    step_header(2, "Tách theo nhân viên", "Mỗi nhân viên sẽ có 1 sheet riêng trong file Excel mới.")
    run = st.button("Tách file", type="primary", key="tach_run")

    if run:
        bar = st.progress(0.0)
        info = st.empty()

        def on_progress(frac, text):
            bar.progress(min(max(frac, 0.0), 1.0))
            info.info(text)

        try:
            res = build_employee_workbook(df, idx_vao, idx_ra, on_progress)
        except Exception as e:
            bar.empty()
            info.empty()
            st.error(f"Có lỗi khi tách file.\n\nChi tiết: {e}")
            return
        bar.empty()
        info.empty()

        if res is None:
            st.warning("Không có nhân viên nào có dữ liệu giờ vào/ra để tách.")
            st.session_state.pop("tach_result", None)
        else:
            data, count_nv, count_tab = res
            st.session_state["tach_result"] = {
                "md5": md5, "bytes": data, "count": count_nv, "tabs": count_tab,
            }
            # Dùng lại ở chức năng 2
            st.session_state["tach_output"] = data

    result = st.session_state.get("tach_result")
    if result and result["md5"] == md5:
        step_header(3, "Tải kết quả")
        stat_cards([
            ("nhân viên đã tách sheet", result["count"], "good"),
            ("sheet có ô trống (tab vàng nhạt)", result["tabs"], "warn" if result["tabs"] else ""),
        ])
        c1, c2 = st.columns([1, 1])
        with c1:
            st.download_button(
                "Tải file Excel tổng hợp",
                result["bytes"],
                "output_tong_hop.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )
        with c2:
            st.button("Tiếp tục: xuất ảnh từ file này", on_click=goto_mode_2, use_container_width=True)
        st.caption("Tải xong nhớ cảm ơn anh Jiang nhé 😄")


# =====================================================================
# CHỨC NĂNG 2: XUẤT BẢNG CÔNG RA ẢNH
# =====================================================================
def build_css(wm_size: int, sp_size: int) -> str:
    return f"""
:root {{
  --bg:#ffffff;
  --bar:#3a3a3a;
  --barText:#ffffff;
  --grid:#b7b7b7;
  --head:#f2f2f2;
  --even:#f7f7f7;
  --total:#1b5e20;
  --excelGreen:#217346;
  --danger:#e60023;
  --dangerBg:#ffe6ea;
}}
*{{box-sizing:border-box;}}
html,body{{margin:0;padding:0;background:var(--bg);}}
body{{
  font-family: Arial, Helvetica, "DejaVu Sans", "Liberation Sans", sans-serif;
  color:#111;
}}
.wrap{{padding:14px 14px 22px;}}
.titlebar{{
  background:var(--bar);
  color:var(--barText);
  text-align:center;
  font-weight:800;
  padding:12px 14px;
  border-radius:8px;
  font-size:22px;
  letter-spacing:.2px;
  margin:0 0 12px 0;
}}
table{{
  border-collapse:collapse;
  width:100%;
  table-layout:auto;
  font-size:14px;
}}
th,td{{
  border:1px solid var(--grid);
  padding:8px 10px;
  text-align:center;
  white-space:nowrap;
}}
thead th{{background:var(--head);font-weight:800;}}
tbody tr:nth-child(even){{background:var(--even);}}

td.highlight{{
  background:var(--excelGreen) !important;
  color:#fff !important;
  font-weight:800;
}}
td.lowhour{{
  background:var(--dangerBg) !important;
  color:var(--danger) !important;
  font-weight:900 !important;
}}
td.lowtext{{
  color:var(--danger) !important;
  font-weight:900 !important;
}}
tr.total-row td{{
  background:var(--total) !important;
  color:#fff !important;
  font-weight:900;
}}
.footer-area{{
  position:relative;
  margin-top:14px;
  min-height:26px;
}}
.stamp{{
  position:absolute;
  left:0;
  bottom:0;
  font-size:{sp_size}px;
  font-weight:800;
  color:#111;
}}
.footer{{
  text-align:center;
  font-size:{wm_size}px;
  color:var(--total);
  font-weight:700;
  line-height:26px;
}}
"""


def safe_sheet_filename(name: str, max_len: int = 90) -> str:
    name = (name or "sheet").strip()
    name = re.sub(r"[\\/:*?\"<>|]+", "_", name)
    name = re.sub(r"\s+", "_", name)
    return name[:max_len] if len(name) > max_len else name


def is_empty_value(v) -> bool:
    if v is None:
        return True
    s = str(v).strip()
    return s == "" or s.lower() in ("nan", "none", "nat")


def parse_float(val) -> float:
    if val is None:
        return 0.0
    s = str(val).strip().replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0


def detect_highlight(cell) -> bool:
    try:
        fill = cell.fill
        if not fill:
            return False
        fg = fill.fgColor
        if not fg:
            return False
        if getattr(fg, "type", None) == "rgb":
            rgb = (fg.rgb or "").upper()
            return rgb not in ("", "00000000", "FFFFFFFF", "FF000000")
        if getattr(fill, "patternType", None) or getattr(fill, "fill_type", None):
            return True
        return False
    except Exception:
        return False


def is_date_header(h: str) -> bool:
    h = (h or "").lower().strip()
    return ("ngày" in h) or ("date" in h)


def is_time_header(h: str) -> bool:
    h = (h or "").lower().strip()
    keys = ["vào", "ra", "gio vao", "gio ra", "giờ vào", "giờ ra", "vao", "ra l", "vào l", "ra lần", "vào lần"]
    return any(k in h for k in keys)


def excel_serial_time_to_hhmm(x: float) -> str:
    total_minutes = int(round(float(x) * 24 * 60))
    hh, mm = divmod(total_minutes, 60)
    hh %= 24
    return f"{hh:02d}:{mm:02d}"


def format_number(x: float) -> str:
    if abs(x - round(x)) < 1e-9:
        return str(int(round(x)))
    return f"{x:.2f}".rstrip("0").rstrip(".")


def format_cell(value, header_str: str) -> str:
    h = header_str or ""
    if value is None:
        return ""

    if isinstance(value, datetime):
        if is_date_header(h):
            return value.strftime("%d/%m/%Y")
        if is_time_header(h):
            return value.strftime("%H:%M")
        return value.strftime("%d/%m/%Y")

    if isinstance(value, date) and not isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")

    if isinstance(value, dtime):
        return f"{value.hour:02d}:{value.minute:02d}"

    if isinstance(value, (float, int)):
        x = float(value)
        if is_time_header(h) and (0 <= x < 1):
            return excel_serial_time_to_hhmm(x)
        return format_number(x)

    return str(value).strip()


def is_col_100(h: str) -> bool:
    s = (h or "").lower()
    return ("lương giờ 100%" in s) or ("luong gio 100%" in s) or ("lương giờ hc" in s) \
        or ("luong gio hc" in s) or ("gio 100%" in s)


def is_col_130(h: str) -> bool:
    s = (h or "").lower()
    return ("lương giờ 130%" in s) or ("luong gio 130%" in s) or ("lương giờ ca đêm" in s) \
        or ("luong gio ca dem" in s) or ("ca đêm" in s) or ("ca dem" in s) \
        or ("tc 130" in s) or ("tăng ca 130" in s)


def normalize_header(x) -> str:
    return str(x).strip() if x is not None else ""


def unique_preserve_order(items):
    seen = set()
    out = []
    for it in items:
        if it not in seen:
            seen.add(it)
            out.append(it)
    return out


def is_total_row(row_vals) -> bool:
    for v in row_vals[: min(6, len(row_vals))]:
        if str(v).strip().lower() in ("tổng", "tong"):
            return True
    return False


def is_text_time_like(s: str) -> bool:
    if not s:
        return False
    return bool(re.match(r"^\d{1,2}:\d{2}$", s.strip()))


def is_numeric_like(v) -> bool:
    if v is None:
        return False
    if isinstance(v, (int, float)):
        try:
            return not (v != v)
        except Exception:
            return True
    if isinstance(v, str):
        s = v.strip()
        if s == "" or is_text_time_like(s):
            return False
        s2 = s.replace(" ", "").replace(",", ".")
        return bool(re.match(r"^-?\d+(\.\d+)?$", s2))
    return False


def should_sum_by_header(h: str) -> bool:
    if not h:
        return True
    hh = h.strip().lower()
    if is_date_header(h) or is_time_header(h):
        return False
    deny_keys = [
        "stt", "thứ", "thu", "ngày", "ngay",
        "mã nv", "ma nv", "họ tên", "ho ten",
        "bộ phận", "bo phan", "ca",
        "ghi chú", "ghi chu", "note", "chú thích", "chu thich",
    ]
    return not any(k in hh for k in deny_keys)


def compute_column_sums(body_rows_raw: list, headers2: list) -> tuple:
    ncol = len(headers2)
    sums = [0.0] * ncol
    counts = [0] * ncol

    for row in body_rows_raw:
        if is_total_row(row):
            continue
        for j in range(ncol):
            if not should_sum_by_header(headers2[j]):
                continue
            v = row[j]
            if is_empty_value(v):
                continue
            if isinstance(v, (datetime, date, dtime)):
                continue
            if isinstance(v, (int, float)):
                sums[j] += float(v)
                counts[j] += 1
            elif isinstance(v, str):
                s = v.strip()
                if s == "" or is_text_time_like(s):
                    continue
                if is_numeric_like(s):
                    sums[j] += parse_float(s)
                    counts[j] += 1
    return sums, counts


def fill_total_row_if_missing(body_rows_raw: list, hl_rows: list, headers2: list, force_overwrite: bool = False):
    sums, counts = compute_column_sums(body_rows_raw, headers2)

    total_idx = None
    for i, row in enumerate(body_rows_raw):
        if is_total_row(row):
            total_idx = i
            break

    ncol = len(headers2)

    if total_idx is None:
        new_total = [None] * ncol
        label_col = 0
        for j, h in enumerate(headers2):
            hs = (h or "").lower()
            if "thứ" in hs or "thu" in hs:
                label_col = j
                break
        new_total[label_col] = "TỔNG"

        for j in range(ncol):
            if counts[j] <= 0 or not should_sum_by_header(headers2[j]):
                continue
            val = sums[j]
            new_total[j] = int(round(val)) if abs(val - round(val)) < 1e-9 else val

        body_rows_raw.append(new_total)
        hl_rows.append([False] * ncol)
        return

    total_row = body_rows_raw[total_idx]
    for j in range(ncol):
        if counts[j] <= 0 or not should_sum_by_header(headers2[j]):
            continue
        need_fill = force_overwrite or is_empty_value(total_row[j])
        if not need_fill:
            continue
        val = sums[j]
        total_row[j] = int(round(val)) if abs(val - round(val)) < 1e-9 else val
    body_rows_raw[total_idx] = total_row


def _norm_key(s: str) -> str:
    s = (s or "").strip().lower()
    return re.sub(r"\s+", " ", s)


def parse_rename_rules(text: str) -> dict:
    mp = {}
    if not text:
        return mp
    for raw in str(text).splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or line.startswith("//"):
            continue
        sep = None
        if "=>" in line:
            sep = "=>"
        elif "->" in line:
            sep = "->"
        elif "=" in line:
            sep = "="
        if not sep:
            continue
        left, right = line.split(sep, 1)
        old, new = left.strip(), right.strip()
        if old and new:
            mp[_norm_key(old)] = new
    return mp


def build_html(sheet_name: str, headers: list, rows: list, stamp_text: str, cfg_css: str,
               watermark: str, show_wm: bool, show_st: bool) -> str:
    ths = "".join(f"<th>{html.escape(str(h) if h is not None else '')}</th>" for h in headers)

    trs = []
    for r in rows:
        tr_cls = "total-row" if r["is_total"] else ""
        tds = []
        for c in r["cells"]:
            classes = []
            if c.get("highlight"):
                classes.append("highlight")
            if c.get("lowhour"):
                classes.append("lowhour")
            if c.get("lowtext"):
                classes.append("lowtext")
            class_attr = f' class="{" ".join(classes)}"' if classes else ""
            tds.append(f"<td{class_attr}>{c['value_html']}</td>")
        trs.append(f'<tr class="{tr_cls}">{"".join(tds)}</tr>')

    stamp_html = f'<div class="stamp">{html.escape(stamp_text)}</div>' if (show_st and stamp_text) else ""
    wm_html = f'<div class="footer">{html.escape(watermark)}</div>' if (show_wm and watermark) else ""

    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>{cfg_css}</style>
</head>
<body>
  <div class="wrap">
    <div class="titlebar">{html.escape(sheet_name)}</div>
    <table>
      <thead><tr>{ths}</tr></thead>
      <tbody>
        {''.join(trs)}
      </tbody>
    </table>
    <div class="footer-area">
      {stamp_html}
      {wm_html}
    </div>
  </div>
</body>
</html>
"""


def find_chromium():
    """Tìm Chromium/Chrome cài sẵn trên hệ thống (ưu tiên hơn bản pyppeteer tự tải)."""
    env_path = os.environ.get("CHROME_PATH") or os.environ.get("CHROMIUM_PATH")
    if env_path and os.path.exists(env_path):
        return env_path
    for name in ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable"):
        found = shutil.which(name)
        if found:
            return found
    for path in ("/usr/bin/chromium", "/usr/bin/chromium-browser", "/usr/bin/google-chrome"):
        if os.path.exists(path):
            return path
    return None


async def render_html_list_to_zip(
    html_list,
    zip_file,
    progress_cb=None,
    dsf: float = 1.8,
    auto_fit: bool = True,
    min_w: int = 980,
    max_w: int = 3500,
    wait_ms: int = 60,
):
    from pyppeteer import launch

    launch_kwargs = dict(
        headless=True,
        handleSIGINT=False,
        handleSIGTERM=False,
        handleSIGHUP=False,
        dumpio=True,  # in log của Chromium ra "Manage app" để dễ tìm lỗi
        args=[
            "--no-sandbox",
            "--disable-setuid-sandbox",
            "--disable-gpu",
            "--disable-dev-shm-usage",
            "--no-zygote",
            "--disable-features=IsolateOrigins,site-per-process",
        ],
    )
    exe = find_chromium()
    if exe:
        launch_kwargs["executablePath"] = exe
    browser = await launch(**launch_kwargs)

    exported = 0
    try:
        page = await browser.newPage()
        await page.setViewport({"width": int(min_w), "height": 900, "deviceScaleFactor": float(dsf)})

        total = len(html_list)
        for i, (fname, html_str) in enumerate(html_list, start=1):
            if progress_cb:
                progress_cb(i - 1, total, fname)

            await page.setContent(html_str)

            if wait_ms > 0:
                try:
                    await page.waitFor(int(wait_ms))
                except Exception:
                    pass

            if auto_fit:
                try:
                    dims = await page.evaluate(
                        "() => ({w: document.documentElement.scrollWidth, h: document.documentElement.scrollHeight})"
                    )
                    w = int(dims.get("w", min_w))
                    w = max(int(min_w), min(int(max_w), w + 6))
                    await page.setViewport({"width": w, "height": 900, "deviceScaleFactor": float(dsf)})
                except Exception:
                    pass

            png_bytes = await page.screenshot({"fullPage": True, "type": "png"})
            zip_file.writestr(fname, png_bytes)
            exported += 1

            if progress_cb:
                progress_cb(i, total, fname)
    finally:
        try:
            await browser.close()
        except Exception:
            pass

    return exported


def run_async(coro):
    loop = asyncio.new_event_loop()
    try:
        asyncio.set_event_loop(loop)
        return loop.run_until_complete(coro)
    finally:
        try:
            loop.close()
        except Exception:
            pass


def _read_ram():
    cur_mb = peak_mb = sys_used_pct = sys_avail_mb = sys_total_mb = None

    if psutil is not None:
        try:
            p = psutil.Process(os.getpid())
            cur_mb = p.memory_info().rss / (1024 * 1024)
            vm = psutil.virtual_memory()
            sys_used_pct = float(vm.percent)
            sys_avail_mb = vm.available / (1024 * 1024)
            sys_total_mb = vm.total / (1024 * 1024)
        except Exception:
            pass

    if resource is not None:
        try:
            ru = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
            peak_mb = (ru / (1024 * 1024)) if sys.platform == "darwin" else (ru / 1024.0)
        except Exception:
            pass

    if peak_mb is None:
        peak_mb = cur_mb

    return cur_mb, peak_mb, sys_used_pct, sys_avail_mb, sys_total_mb


def hdr_key(h: str) -> str:
    return f"hdr__{hashlib.md5(h.encode('utf-8')).hexdigest()[:10]}"


def _toggle_header(h: str):
    s = set(st.session_state.get("selected_headers") or set())
    if st.session_state.get(hdr_key(h)):
        s.add(h)
    else:
        s.discard(h)
    st.session_state.selected_headers = s


def _set_all_headers(headers, value: bool):
    st.session_state.selected_headers = set(headers) if value else set()
    for h in headers:
        st.session_state[hdr_key(h)] = value


def prepare_sheets(wb, selected_headers, opt, on_progress=None):
    """Tạo danh sách (tên ảnh, HTML) cho từng sheet. Trả về (to_render, prepared, skipped, total_sheets)."""
    rename_map = opt["rename_map"]

    def rename_header(h):
        return rename_map.get(_norm_key(h), h) if rename_map else h

    cfg_css = build_css(opt["watermark_size"], opt["stamp_size"])
    to_render = []
    sheetnames = wb.sheetnames
    total_sheets = len(sheetnames)
    prepared = 0
    skipped = 0

    for idx_sheet, sheet_name in enumerate(sheetnames, start=1):
        if on_progress:
            on_progress(idx_sheet / max(total_sheets, 1), f"Đang chuẩn bị sheet {idx_sheet}/{total_sheets}: {sheet_name}")
        ws = wb[sheet_name]

        data = []
        hl = []
        max_cols = 0

        for row in ws.iter_rows():
            row_vals = []
            row_hls = []
            for cell in row:
                row_vals.append(cell.value)
                row_hls.append(detect_highlight(cell))
            max_cols = max(max_cols, len(row_vals))
            data.append(row_vals)
            hl.append(row_hls)

        if len(data) < 2:
            skipped += 1
            continue

        for r in range(len(data)):
            if len(data[r]) < max_cols:
                data[r].extend([None] * (max_cols - len(data[r])))
            if len(hl[r]) < max_cols:
                hl[r].extend([False] * (max_cols - len(hl[r])))

        keep_cols = [j for j in range(max_cols)
                     if any(not is_empty_value(data[i][j]) for i in range(1, len(data)))]
        if not keep_cols:
            skipped += 1
            continue

        headers_full = [normalize_header(data[0][j]) for j in keep_cols]
        keep_cols2 = []
        headers2 = []
        for j, hname in zip(keep_cols, headers_full):
            if hname and (hname in selected_headers):
                keep_cols2.append(j)
                headers2.append(hname)

        if not keep_cols2:
            skipped += 1
            continue

        body_rows_raw = [[data[i][j] for j in keep_cols2] for i in range(1, len(data))]
        hl_rows = [[hl[i][j] for j in keep_cols2] for i in range(1, len(hl))]
        fill_total_row_if_missing(body_rows_raw, hl_rows, headers2, force_overwrite=False)

        col_100 = None
        col_130 = None
        for i, h in enumerate(headers2):
            if col_100 is None and is_col_100(h):
                col_100 = i
            if col_130 is None and is_col_130(h):
                col_130 = i

        headers_display = [rename_header(h) for h in headers2]
        headers_out = (["STT"] + headers_display) if opt["add_stt"] else headers_display

        rows_for_html = []
        stt_counter = 0
        threshold = float(opt["threshold"])

        for r_idx, row in enumerate(body_rows_raw):
            is_total = is_total_row(row)

            stt_val = ""
            if opt["add_stt"] and (not is_total):
                stt_counter += 1
                stt_val = str(stt_counter)

            low_bg_cols = set()
            low_text_cols = set()

            if (not is_total) and (col_100 is not None) and (col_130 is not None):
                raw100 = row[col_100]
                raw130 = row[col_130]
                if (not is_empty_value(raw100)) or (not is_empty_value(raw130)):
                    if (parse_float(raw100) + parse_float(raw130)) < threshold:
                        low_bg_cols.add(col_100)
                        low_bg_cols.add(col_130)
            elif (not is_total) and (col_100 is not None) and (col_130 is None):
                raw100 = row[col_100]
                if not is_empty_value(raw100) and parse_float(raw100) < threshold:
                    low_text_cols.add(col_100)

            cells = []
            if opt["add_stt"]:
                cells.append({
                    "value_html": html.escape(stt_val),
                    "highlight": False, "lowhour": False, "lowtext": False,
                })

            for c_idx, val in enumerate(row):
                formatted = format_cell(val, headers2[c_idx])
                highlight = bool(hl_rows[r_idx][c_idx]) if r_idx < len(hl_rows) and c_idx < len(hl_rows[r_idx]) else False
                cells.append({
                    "value_html": html.escape(formatted),
                    "highlight": highlight,
                    "lowhour": (c_idx in low_bg_cols),
                    "lowtext": (c_idx in low_text_cols),
                })

            rows_for_html.append({"is_total": is_total, "cells": cells})

        base_name = safe_sheet_filename(sheet_name)
        fname = f"{prepared + 1}_{base_name}.png"

        if opt["stamp_mode"] == "Tên sheet":
            stamp_text = base_name
        elif opt["stamp_mode"] == "Tên file ảnh":
            stamp_text = fname
        else:
            stamp_text = (opt["stamp_custom"] or "").strip()

        html_doc = build_html(
            sheet_name=sheet_name,
            headers=headers_out,
            rows=rows_for_html,
            stamp_text=stamp_text,
            cfg_css=cfg_css,
            watermark=opt["watermark_text"],
            show_wm=opt["show_watermark"],
            show_st=opt["show_stamp"],
        )

        to_render.append((fname, html_doc))
        prepared += 1

    return to_render, prepared, skipped, total_sheets


def page_xuat_anh():
    # ---------- Bước 1: chọn nguồn file ----------
    step_header(1, "Chọn file Excel tổng hợp", "Là file kết quả của chức năng “Tách file chấm công”.")

    source = "Tải file lên"
    if "tach_output" in st.session_state:
        source = st.radio(
            "Nguồn file",
            ["Dùng file vừa tách", "Tải file lên"],
            horizontal=True,
            label_visibility="collapsed",
            key="anh_source",
        )

    if source == "Tải file lên":
        uploaded = st.file_uploader(
            "File Excel tổng hợp", type=["xlsx"], key="anh_uploader", label_visibility="collapsed"
        )
        if not uploaded:
            st.info("Chưa có file. Kéo thả file Excel tổng hợp vào khung phía trên để bắt đầu.")
            return
        file_bytes = uploaded.getvalue()
    else:
        file_bytes = st.session_state["tach_output"]
        st.success("Đang dùng file tổng hợp vừa tách ở chức năng 1.")

    file_md5 = hashlib.md5(file_bytes).hexdigest()

    if st.session_state.get("file_md5") != file_md5:
        try:
            with st.spinner("Đang đọc file..."):
                wb_loaded = load_workbook(BytesIO(file_bytes), data_only=True)
        except Exception as e:
            st.error(f"Không đọc được file. Hãy kiểm tra đây có phải file .xlsx hợp lệ không.\n\nChi tiết: {e}")
            return
        st.session_state.file_md5 = file_md5
        st.session_state.wb = wb_loaded
        st.session_state.all_headers = None
        st.session_state.selected_headers = None
        st.session_state.pop("anh_result", None)
        for k in list(st.session_state.keys()):
            if isinstance(k, str) and k.startswith("hdr__"):
                del st.session_state[k]

    wb = st.session_state.wb

    if st.session_state.get("all_headers") is None:
        tmp = []
        for sn in wb.sheetnames:
            ws = wb[sn]
            first_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))
            if not first_row:
                continue
            headers = [normalize_header(x) for x in list(first_row[0])]
            tmp.extend([h for h in headers if h != ""])
        st.session_state.all_headers = unique_preserve_order(tmp)

    if st.session_state.get("selected_headers") is None:
        st.session_state.selected_headers = set(st.session_state.all_headers)

    all_headers = st.session_state.all_headers
    if not all_headers:
        st.warning("Không tìm thấy dòng tiêu đề trong file. Dòng đầu tiên của mỗi sheet cần là tên các cột.")
        return

    stat_cards([("sheet trong file", len(wb.sheetnames), ""), ("cột khác nhau", len(all_headers), "")])

    # ---------- Bước 2: chọn cột ----------
    step_header(2, "Chọn cột hiển thị trên ảnh", "Tick cột nào thì ảnh giữ cột đó. Áp dụng cho tất cả sheet.")

    with st.container(border=True):
        cA, cB, cC = st.columns([1, 1, 3])
        cA.button("Chọn tất cả", on_click=_set_all_headers, args=(all_headers, True), use_container_width=True)
        cB.button("Bỏ chọn tất cả", on_click=_set_all_headers, args=(all_headers, False), use_container_width=True)
        search = cC.text_input(
            "Tìm cột", value="", placeholder="Tìm cột… ví dụ: lương, vào, ra, tăng ca",
            label_visibility="collapsed",
        ).strip().lower()

        filtered = [h for h in all_headers if (search in h.lower())] if search else all_headers

        grid = st.columns(3)
        master = set(st.session_state.selected_headers)
        for i, h in enumerate(filtered):
            key = hdr_key(h)
            if key not in st.session_state:
                st.session_state[key] = (h in master)
            grid[i % 3].checkbox(h, key=key, on_change=_toggle_header, args=(h,))
        if search and not filtered:
            st.caption("Không có cột nào khớp từ khóa.")

    selected_headers = set(st.session_state.selected_headers)
    st.caption(f"Đã chọn {len(selected_headers)}/{len(all_headers)} cột")

    # ---------- Bước 3: tùy chỉnh ----------
    step_header(3, "Tùy chỉnh ảnh", "Không bắt buộc. Có thể bỏ qua nếu dùng cấu hình mặc định.")

    with st.expander("Mở tùy chỉnh"):
        t1, t2, t3, t4 = st.tabs(["Chất lượng ảnh", "Bảng công", "Watermark & dấu", "Đổi tên cột"])
        with t1:
            device_scale = st.slider("Độ nét (càng cao càng nét nhưng chạy chậm)", 1.0, 3.0, 1.8, 0.1)
            wait_ms = st.slider("Chờ trang hiển thị xong (ms)", 0, 300, 60, 10)
            auto_fit_width = st.checkbox("Tự co giãn chiều ngang cho vừa bảng", value=True)
            w1, w2 = st.columns(2)
            min_width = w1.number_input("Chiều ngang tối thiểu (px)", min_value=600, max_value=6000, value=980, step=10)
            max_width = w2.number_input("Chiều ngang tối đa (px)", min_value=800, max_value=12000, value=3500, step=10)
            st.caption("Gợi ý: nếu chạy chậm, giảm độ nét xuống 1.4–1.8 và bỏ bớt cột.")
        with t2:
            add_stt = st.checkbox("Thêm cột STT", value=True)
            threshold_hours = st.number_input(
                "Ngưỡng giờ chuẩn (nhỏ hơn ngưỡng thì tô đỏ)",
                min_value=0.0, max_value=24.0, value=float(DEFAULT_THRESHOLD), step=0.5,
            )
        with t3:
            show_watermark = st.checkbox("Hiện watermark (giữa, phía dưới)", value=True)
            watermark_text = st.text_input("Nội dung watermark", value=DEFAULT_WATERMARK)
            watermark_size = st.slider("Cỡ chữ watermark", 10, 30, 18)
            st.divider()
            show_stamp = st.checkbox("Hiện dấu (góc dưới bên trái)", value=True)
            stamp_mode = st.selectbox("Nội dung dấu", ["Tên sheet", "Tên file ảnh", "Tùy chỉnh"], index=0)
            stamp_custom = st.text_input("Nội dung dấu tùy chỉnh", value="", disabled=(stamp_mode != "Tùy chỉnh"))
            stamp_size = st.slider("Cỡ chữ dấu", 10, 26, 16)
        with t4:
            rename_enable = st.checkbox("Đổi tên tiêu đề cột khi xuất ảnh", value=True)
            rename_rules = st.text_area(
                "Mỗi dòng 1 cặp: Tên cũ = Tên mới",
                value=DEFAULT_RENAME_RULES,
                height=170,
                disabled=(not rename_enable),
            )

    opt = {
        "rename_map": parse_rename_rules(rename_rules) if rename_enable else {},
        "watermark_size": watermark_size,
        "stamp_size": stamp_size,
        "add_stt": add_stt,
        "threshold": threshold_hours,
        "stamp_mode": stamp_mode,
        "stamp_custom": stamp_custom,
        "watermark_text": watermark_text,
        "show_watermark": show_watermark,
        "show_stamp": show_stamp,
    }

    # ---------- Bước 4: xuất ảnh ----------
    step_header(4, "Xuất ảnh", "Mỗi sheet thành 1 ảnh PNG, nén chung trong 1 file ZIP.")

    mem_box = st.expander("Thông tin hệ thống (RAM)")
    mem_placeholder = mem_box.empty()
    mem_state = {"peak": 0.0, "t0": time.perf_counter()}

    def render_mem_ui(cur_mb, peak_mb, sys_used_pct, sys_avail_mb, sys_total_mb):
        lines = []
        lines.append(f"**RAM tiến trình (hiện tại):** `{cur_mb:.0f} MB`" if cur_mb is not None
                     else "**RAM tiến trình (hiện tại):** `—`")
        if peak_mb is not None:
            lines.append(f"**RAM tiến trình (cao nhất):** `{peak_mb:.0f} MB`")
        if sys_used_pct is not None and sys_avail_mb is not None and sys_total_mb is not None:
            lines.append(f"**RAM hệ thống:** `{sys_used_pct:.0f}%` — còn `{sys_avail_mb:.0f}/{sys_total_mb:.0f} MB`")
        lines.append(f"**Thời gian chạy:** `{time.perf_counter() - mem_state['t0']:.1f}s`")
        mem_placeholder.markdown("\n\n".join(lines))

    # Hiện RAM ngay khi mở, không đợi đến lúc xuất ảnh
    _cur, _peak, _pct, _avail, _total = _read_ram()
    if _cur is not None:
        mem_state["peak"] = max(mem_state["peak"], _cur)
    render_mem_ui(_cur, mem_state["peak"] if _cur is not None else _peak, _pct, _avail, _total)

    render_clicked = st.button("Xuất ảnh", type="primary", key="anh_run")

    if render_clicked:
        if not selected_headers:
            st.error("Bạn đang bỏ chọn hết cột. Hãy chọn ít nhất 1 cột ở bước 2.")
            return

        bar = st.progress(0.0)
        status = st.empty()

        def on_prep(frac, text):
            bar.progress(min(max(frac, 0.0), 1.0) * 0.3)
            status.info(text)

        to_render, prepared, skipped, total_sheets = prepare_sheets(wb, selected_headers, opt, on_prep)

        if not to_render:
            bar.empty()
            status.empty()
            st.warning("Không có sheet nào chứa các cột đã chọn. Hãy chọn thêm cột ở bước 2.")
            return

        def progress_cb(done, total, current_name):
            done = max(0, min(done, total))
            pct = 0.0 if total == 0 else done / total
            bar.progress(0.3 + 0.7 * pct)
            status.info(f"Đang xuất ảnh {done}/{total}: {current_name}")

            if (done == 0) or (done == total) or (done % 3 == 0):
                cur_mb, peak_mb, sys_used_pct, sys_avail_mb, sys_total_mb = _read_ram()
                if cur_mb is not None:
                    mem_state["peak"] = max(mem_state["peak"], cur_mb)
                if peak_mb is not None:
                    mem_state["peak"] = max(mem_state["peak"], peak_mb)
                render_mem_ui(cur_mb, mem_state["peak"], sys_used_pct, sys_avail_mb, sys_total_mb)

        spool = SpooledTemporaryFile(max_size=64 * 1024 * 1024)
        exported_count = 0
        try:
            with zipfile.ZipFile(spool, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
                exported_count = run_async(
                    render_html_list_to_zip(
                        to_render,
                        z,
                        progress_cb=progress_cb,
                        dsf=device_scale,
                        auto_fit=auto_fit_width,
                        min_w=int(min_width),
                        max_w=int(max_width),
                        wait_ms=int(wait_ms),
                    )
                )
        except Exception as e:
            bar.empty()
            status.empty()
            st.error(f"Không xuất được ảnh (lỗi trình duyệt Chromium).\n\nChi tiết: {e}\n\nKiểm tra file packages.txt có dòng `chromium` và xem log trong Manage app.")
            return

        bar.empty()
        status.empty()

        spool.seek(0)
        zip_bytes = spool.read()
        st.session_state["anh_result"] = {
            "md5": file_md5,
            "zip": zip_bytes,
            "exported": exported_count,
            "prepared": prepared,
            "skipped": skipped,
            "total": total_sheets,
        }

    res = st.session_state.get("anh_result")
    if res and res["md5"] == file_md5:
        st.success("Đã xuất xong. Tải file ZIP bên dưới.")
        stat_cards([
            ("ảnh đã xuất", res["exported"], "good"),
            ("sheet bị bỏ qua", res["skipped"], "warn" if res["skipped"] else ""),
            ("dung lượng ZIP", f"{len(res['zip']) / (1024 * 1024):.1f} MB", ""),
        ])
        st.download_button(
            "Tải ZIP ảnh PNG",
            data=res["zip"],
            file_name="bang_cong_png.zip",
            mime="application/zip",
            type="primary",
        )


# =====================================================================
# ĐĂNG NHẬP BẰNG MẬT KHẨU
# =====================================================================
MAX_FAILS = 5
LOCK_SECONDS = 60


def require_login():
    """Chặn toàn bộ app cho đến khi nhập đúng mật khẩu (lưu trong Streamlit Secrets: APP_PASSWORD)."""
    if st.session_state.get("auth_ok"):
        return

    try:
        expected = str(st.secrets["APP_PASSWORD"])
    except Exception:
        expected = ""

    if not expected:
        st.error(
            "Chưa cấu hình mật khẩu cho app. Vào Streamlit Cloud → Settings → Secrets "
            "và thêm dòng: APP_PASSWORD = \"mật-khẩu-của-bạn\""
        )
        st.stop()

    locked_until = st.session_state.get("lock_until", 0.0)
    now = time.time()
    if now < locked_until:
        st.error(f"Nhập sai quá nhiều lần. Thử lại sau {int(locked_until - now) + 1} giây.")
        st.stop()

    step_header(1, "Đăng nhập", "Nhập mật khẩu để sử dụng công cụ.")
    with st.form("login_form"):
        pw = st.text_input("Mật khẩu", type="password", label_visibility="collapsed", placeholder="Mật khẩu")
        submitted = st.form_submit_button("Đăng nhập", type="primary")

    # Gợi ý mật khẩu (tùy chọn): đặt APP_HINT trong Streamlit Secrets, không ghi trong code
    try:
        hint = str(st.secrets.get("APP_HINT", "")).strip()
    except Exception:
        hint = ""
    if hint:
        st.markdown(
            f'<div style="margin-top:.5rem;color:#5B6B73;font-size:.92rem;">Gợi ý mật khẩu: '
            f'<b>{html.escape(hint)}</b></div>',
            unsafe_allow_html=True,
        )

    if submitted:
        if hmac.compare_digest(pw.encode("utf-8"), expected.encode("utf-8")):
            st.session_state["auth_ok"] = True
            st.session_state["auth_fail"] = 0
            st.rerun()
        else:
            fails = st.session_state.get("auth_fail", 0) + 1
            st.session_state["auth_fail"] = fails
            if fails >= MAX_FAILS:
                st.session_state["lock_until"] = time.time() + LOCK_SECONDS
                st.session_state["auth_fail"] = 0
                st.error(f"Nhập sai {MAX_FAILS} lần. Đã khóa {LOCK_SECONDS} giây.")
            else:
                st.error(f"Mật khẩu chưa đúng. Còn {MAX_FAILS - fails} lần thử.")
    st.stop()


def logout():
    st.session_state["auth_ok"] = False


# =====================================================================
# MAIN
# =====================================================================
inject_css()
hero()
require_login()

if "mode" not in st.session_state:
    st.session_state["mode"] = MODE_1

st.radio("Chức năng", [MODE_1, MODE_2], key="mode", horizontal=True, label_visibility="collapsed")
st.markdown(f'<div class="mode-hint">{html.escape(MODE_HINT[st.session_state["mode"]])}</div>', unsafe_allow_html=True)

if st.session_state["mode"] == MODE_1:
    page_tach_file()
else:
    page_xuat_anh()

st.markdown('<div class="footer-note">Production by Jiangvux</div>', unsafe_allow_html=True)
st.button("Đăng xuất", on_click=logout, key="logout_btn")
